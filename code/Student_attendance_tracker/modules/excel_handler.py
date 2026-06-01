
import openpyxl

from config.settings import EXCEL_FILE

book = openpyxl.load_workbook(EXCEL_FILE)

sheet = book.active


def savefile():

    book.save(EXCEL_FILE)

    print("saved!")


def get_student(row):

    roll = sheet.cell(row=row, column=1).value

    name = sheet.cell(row=row, column=2).value

    email = sheet.cell(row=row, column=3).value

    return roll, name, email


def get_leave(row, col):

    return sheet.cell(row=row, column=col).value


def update_leave(row, col, value):

    sheet.cell(row=row, column=col).value = value